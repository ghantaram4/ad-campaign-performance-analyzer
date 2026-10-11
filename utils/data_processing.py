"""Load and validate campaign data without silently replacing invalid values."""
from io import BytesIO, StringIO
import csv
from zipfile import ZipFile
from defusedxml.ElementTree import iterparse
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
from pathlib import Path
import numpy as np
import pandas as pd

REQUIRED_COLUMNS = [
    'Campaign_ID', 'Target_Audience', 'Campaign_Goal', 'Duration', 'Channel_Used',
    'Conversion_Rate', 'Acquisition_Cost', 'ROI', 'Location', 'Language', 'Clicks',
    'Impressions', 'Engagement_Score', 'Customer_Segment', 'Date', 'Company',
]
CATEGORIES = ['Target_Audience', 'Campaign_Goal', 'Channel_Used', 'Customer_Segment']

class DataValidationError(ValueError):
    """A readable input error for the dashboard."""

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10_000
MAX_COLUMNS = 64
MAX_CELL_CHARS = 2048
MAX_XLSX_EXPANDED_BYTES = 40 * 1024 * 1024
MAX_ZIP_ENTRIES = 256
MAX_COMPRESSION_RATIO = 200
MAX_NUMERIC_VALUE = 1_000_000_000_000


def validate_headers(headers):
    names = [str(value).strip() if value is not None else '' for value in headers]
    if not names or len(names) > MAX_COLUMNS:
        raise DataValidationError('Files must contain at most 64 columns.')
    if len(set(names)) != len(names):
        raise DataValidationError('Column names must be unique.')
    missing = sorted(set(REQUIRED_COLUMNS) - set(names))
    if missing:
        raise DataValidationError('Missing required columns: ' + ', '.join(missing))
    return names


def check_cells(values):
    if len(values) > MAX_COLUMNS or any(isinstance(v, str) and len(v) > MAX_CELL_CHARS for v in values):
        raise DataValidationError('Files must have at most 64 columns and 2,048 characters per cell.')
    if any(isinstance(v, str) and '\x00' in v for v in values):
        raise DataValidationError('Text cells cannot contain null characters.')


def validate_xlsx_archive(content):
    # No extraction to disk. Check both metadata and actual expanded bytes before openpyxl.
    with ZipFile(BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_ZIP_ENTRIES or len({e.filename for e in entries}) != len(entries):
            raise DataValidationError('Workbook has too many or duplicate archive entries.')
        names = {e.filename for e in entries}
        if not {'[Content_Types].xml', 'xl/workbook.xml'} <= names:
            raise DataValidationError('The file is not a valid XLSX workbook.')
        if any(e.flag_bits & 1 or e.file_size > MAX_XLSX_EXPANDED_BYTES or
               e.file_size / max(e.compress_size, 1) > MAX_COMPRESSION_RATIO for e in entries):
            raise DataValidationError('Workbook is encrypted or exceeds safe decompression limits.')
        if sum(e.file_size for e in entries) > MAX_XLSX_EXPANDED_BYTES:
            raise DataValidationError('Workbook expanded size exceeds 40 MB.')
        expanded = 0
        for entry in entries:
            if 'vbaproject' in entry.filename.lower() or entry.filename.startswith('xl/externalLinks/'):
                raise DataValidationError('Macros and external workbook links are not supported.')
            with archive.open(entry) as stream:
                chunks = []
                size = 0
                while chunk := stream.read(64 * 1024):
                    size += len(chunk)
                    expanded += len(chunk)
                    if expanded > MAX_XLSX_EXPANDED_BYTES or size > entry.file_size:
                        raise DataValidationError('Workbook exceeds safe decompression limits.')
                    if entry.filename.endswith(('.xml', '.rels')):
                        chunks.append(chunk)
                if chunks:
                    # Defused XML rejects entities/DTD; free each completed node.
                    shared_strings = 0
                    worksheet = entry.filename.startswith('xl/worksheets/')
                    for _, element in iterparse(BytesIO(b''.join(chunks)), events=('end',), forbid_dtd=True):
                        tag = element.tag.rsplit('}', 1)[-1]
                        if worksheet and tag == 'row' and int(element.get('r', '0')) > MAX_ROWS + 1:
                            raise DataValidationError('Files must contain at most 10,000 campaign rows.')
                        if worksheet and tag == 'c':
                            # Reject sparse wide cells before openpyxl can expand padding.
                            coordinate = element.get('r')
                            if coordinate is None or column_index_from_string(coordinate_from_string(coordinate)[0]) > MAX_COLUMNS:
                                raise DataValidationError('Files must contain at most 64 columns.')
                        if worksheet and tag == 'f':
                            raise DataValidationError('Excel formulas are not supported. Export values before uploading.')
                        if tag == 't' and element.text and len(element.text) > MAX_CELL_CHARS:
                            raise DataValidationError('Files must have at most 2,048 characters per cell.')
                        if entry.filename == 'xl/sharedStrings.xml' and tag == 'si':
                            shared_strings += 1
                            if shared_strings > MAX_ROWS * MAX_COLUMNS:
                                raise DataValidationError('Workbook contains too many shared strings.')
                        element.clear()


def read_data(content: bytes, filename: str) -> pd.DataFrame:
    if len(content) > MAX_UPLOAD_BYTES:
        raise DataValidationError('Upload exceeds the 10 MB limit.')
    extension = Path(filename).suffix.lower()
    if extension not in {'.csv', '.xlsx'}:
        raise DataValidationError('Upload a UTF-8 CSV or an .xlsx workbook.')
    try:
        rows = []
        if extension == '.xlsx':
            validate_xlsx_archive(content)
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.worksheets[0]
                if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                    raise DataValidationError('Files must contain at most 64 columns.')
                if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                    raise DataValidationError('Files must contain at most 10,000 campaign rows.')
                # Do not trust worksheet dimensions to hide or omit actual cells.
                sheet.reset_dimensions()
                iterator = sheet.iter_rows()
                first = next(iterator, ())
                if any(cell.data_type == 'f' for cell in first):
                    raise DataValidationError('Excel formulas are not supported. Export values before uploading.')
                headers = validate_headers([cell.value for cell in first])
                check_cells(headers)
                for index, cells in enumerate(iterator):
                    if index >= MAX_ROWS:
                        raise DataValidationError('Files must contain at most 10,000 campaign rows.')
                    if any(cell.data_type == 'f' for cell in cells):
                        raise DataValidationError('Excel formulas are not supported. Export values before uploading.')
                    values = [cell.value for cell in cells]
                    check_cells(values)
                    if len(values) > len(headers):
                        raise DataValidationError('A row contains more cells than the header.')
                    rows.append(values + [None] * (len(headers) - len(values)))
            finally:
                workbook.close()
        else:
            reader = csv.reader(StringIO(content.decode('utf-8-sig'), newline=''), strict=True)
            headers = validate_headers(next(reader, []))
            check_cells(headers)
            for index, values in enumerate(reader):
                if index >= MAX_ROWS:
                    raise DataValidationError('Files must contain at most 10,000 campaign rows.')
                if not values:
                    continue
                check_cells(values)
                if len(values) != len(headers):
                    raise DataValidationError('Every CSV row must have the same number of fields as the header.')
                rows.append([value if value != '' else None for value in values])
        if extension == '.csv':
            # Restore pandas type inference only after structural/resource validation.
            return pd.read_csv(BytesIO(content), header=0, names=headers,
                               dtype={'Campaign_ID': str}, encoding='utf-8-sig', nrows=MAX_ROWS + 1)
        frame = pd.DataFrame(rows, columns=headers)
        frame['Campaign_ID'] = frame['Campaign_ID'].map(lambda v: str(v) if v is not None else None)
        return frame
    except DataValidationError:
        raise
    except Exception:
        # Parser exceptions may contain cell values; do not expose or chain them.
        raise DataValidationError('Cannot read this file. Upload a UTF-8 CSV or an .xlsx workbook with headers in the first row.') from None


def clean_data(raw: pd.DataFrame) -> pd.DataFrame:
    if len(raw) > MAX_ROWS:
        raise DataValidationError('Files must contain at most 10,000 campaign rows.')
    if len(raw.columns) > MAX_COLUMNS:
        raise DataValidationError('Files must contain at most 64 columns.')
    for values in raw.itertuples(index=False, name=None):
        check_cells(values)
    df = raw.copy()
    df.columns = df.columns.astype(str).str.strip()
    if df.columns.duplicated().any():
        raise DataValidationError('Column names must be unique.')
    missing = sorted(set(REQUIRED_COLUMNS) - set(df.columns))
    if missing:
        raise DataValidationError('Missing required columns: ' + ', '.join(missing))
    if df.empty:
        raise DataValidationError('The uploaded file contains headers but no campaign rows.')
    df = df[REQUIRED_COLUMNS].copy()
    for col in ['Campaign_ID', 'Company', 'Location', 'Language'] + CATEGORIES:
        if df[col].isna().any() or df[col].astype(str).str.strip().eq('').any():
            raise DataValidationError(f'{col} contains blank values.')
        df[col] = df[col].astype(str).str.strip()
        if col in CATEGORIES and df[col].nunique() > 100:
            raise DataValidationError(f'{col} must contain at most 100 distinct categories.')
    if df.Campaign_ID.duplicated().any():
        raise DataValidationError('Campaign_ID must be unique: this dashboard expects one row per campaign.')
    for col in ['Clicks', 'Impressions', 'Engagement_Score', 'ROI', 'Conversion_Rate']:
        values = df[col].astype(str).str.strip()
        if col in ['ROI', 'Conversion_Rate']:
            percent = values.str.endswith('%')
            values = values.str.replace('%', '', regex=False)
        df[col] = pd.to_numeric(values, errors='coerce')
        if col == 'Conversion_Rate':
            df.loc[percent, col] /= 100
        if not np.isfinite(df[col].to_numpy(dtype=float)).all() or (df[col].abs() > MAX_NUMERIC_VALUE).any():
            raise DataValidationError(f'{col} contains a missing or invalid numeric value.')
    df['Spend'] = pd.to_numeric(
        df.Acquisition_Cost.astype(str).str.strip().str.replace(r'[$,]', '', regex=True), errors='coerce')
    if not np.isfinite(df.Spend.to_numpy(dtype=float)).all() or (df.Spend < 0).any() or (df.Spend > MAX_NUMERIC_VALUE).any():
        raise DataValidationError('Acquisition_Cost must contain non-negative USD amounts, such as $420.00.')
    for col in ['Clicks', 'Impressions']:
        if (df[col] < 0).any() or (df[col] % 1 != 0).any():
            raise DataValidationError(f'{col} must contain non-negative whole counts.')
        df[col] = df[col].astype('int64')
    if (df.Clicks > df.Impressions).any():
        raise DataValidationError('Clicks cannot exceed impressions.')
    if not df.Conversion_Rate.between(0, 1).all():
        raise DataValidationError('Conversion_Rate must be a decimal between 0 and 1, or a percentage such as 12%.')
    # ISO strings and Excel-native datetimes are supported. Reject ambiguous local dates.
    date_text = df.Date.astype(str).str.strip()
    iso = date_text.str.match(r'^\d{4}-\d{2}-\d{2}(?:[ T].*)?$')
    if not iso.all():
        raise DataValidationError('Date must use YYYY-MM-DD (or a native Excel date). Widening an Excel column does not alter its underlying dates.')
    df['Date'] = pd.to_datetime(date_text, format='mixed', errors='coerce')
    if df.Date.isna().any():
        raise DataValidationError('Date contains invalid dates. Use YYYY-MM-DD.')
    try:
        df['Date'] = df.Date.dt.normalize()
    except AttributeError as exc:
        raise DataValidationError('Use dates without mixed timezones.') from exc
    if df.Date.dt.tz is not None:
        raise DataValidationError('Use dates without timezone offsets.')
    duration = df.Duration.astype(str).str.extract(r'^\s*(\d+)\s*(?:Days?)?\s*$', expand=False)
    df['Duration_Days'] = pd.to_numeric(duration, errors='coerce')
    if df.Duration_Days.isna().any() or (df.Duration_Days <= 0).any() or (df.Duration_Days > 365000).any():
        raise DataValidationError('Duration must be positive whole days, such as 15 Days.')
    df['Duration_Days'] = df.Duration_Days.astype('int64')
    df['Conversions'] = (df.Clicks * df.Conversion_Rate).round().astype('int64')
    return df.sort_values(['Date', 'Campaign_ID']).reset_index(drop=True)

def apply_filters(df, start, end, channels, audiences, goals, segments):
    mask = df.Date.between(pd.Timestamp(start), pd.Timestamp(end))
    for col, selected in [('Channel_Used', channels), ('Target_Audience', audiences),
                          ('Campaign_Goal', goals), ('Customer_Segment', segments)]:
        mask &= df[col].isin(selected)
    return df.loc[mask].copy()

def search_campaigns(df, query):
    if not query.strip():
        return df.copy()
    searchable = df[['Campaign_ID', 'Company', 'Channel_Used', 'Campaign_Goal', 'Target_Audience', 'Customer_Segment']]
    mask = searchable.apply(lambda col: col.str.contains(query.strip(), case=False, regex=False)).any(axis=1)
    return df.loc[mask].copy()
